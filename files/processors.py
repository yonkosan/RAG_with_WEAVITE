import os
import io
import magic
import PyPDF2
import pdfplumber
from docx import Document as DocxDocument
from PIL import Image
import pytesseract
import pandas as pd
from typing import List, Dict, Any, Optional
import json
import cv2
import numpy as np

class FileProcessor:
    """Base class for file processing."""
    
    SUPPORTED_EXTENSIONS = {
        'pdf', 'docx', 'doc', 'txt', 'csv', 'xlsx', 'xls', 
        'jpg', 'jpeg', 'png', 'bmp', 'tiff', 'json'
    }
    
    @classmethod
    def get_file_type(cls, file_path: str) -> str:
        """Detect file type using python-magic."""
        try:
            mime = magic.from_file(file_path, mime=True)
            return mime
        except Exception:
            # Fallback to extension-based detection
            ext = os.path.splitext(file_path)[1].lower().lstrip('.')
            return f"application/{ext}" if ext in cls.SUPPORTED_EXTENSIONS else "unknown"
    
    @classmethod
    def is_supported(cls, filename: str) -> bool:
        """Check if file type is supported."""
        ext = os.path.splitext(filename)[1].lower().lstrip('.')
        return ext in cls.SUPPORTED_EXTENSIONS

class PDFProcessor:
    """Process PDF files to extract text, images, and tables."""
    
    @staticmethod
    def extract_text(file_path: str) -> str:
        """Extract text from PDF using PyPDF2."""
        try:
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return text.strip()
        except Exception as e:
            print(f"Error extracting text from PDF: {e}")
            return ""
    
    @staticmethod
    def extract_tables(file_path: str) -> List[Dict[str, Any]]:
        """Extract tables from PDF using pdfplumber."""
        tables = []
        try:
            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages):
                    page_tables = page.extract_tables()
                    for table_num, table in enumerate(page_tables):
                        if table and len(table) > 1:  # Ensure table has header and data
                            df = pd.DataFrame(table[1:], columns=table[0])
                            tables.append({
                                'page': page_num + 1,
                                'table_number': table_num + 1,
                                'data': df.to_dict('records'),
                                'csv': df.to_csv(index=False)
                            })
        except Exception as e:
            print(f"Error extracting tables from PDF: {e}")
        return tables
    
    @staticmethod
    def extract_comprehensive(file_path: str) -> Dict[str, Any]:
        """Extract text and tables from PDF."""
        return {
            'text': PDFProcessor.extract_text(file_path),
            'tables': PDFProcessor.extract_tables(file_path)
        }

class DocxProcessor:
    """Process DOCX files to extract text and tables."""
    
    @staticmethod
    def extract_text(file_path: str) -> str:
        """Extract text from DOCX file."""
        try:
            doc = DocxDocument(file_path)
            paragraphs = [para.text for para in doc.paragraphs]
            return "\n".join(paragraphs)
        except Exception as e:
            print(f"Error extracting text from DOCX: {e}")
            return ""
    
    @staticmethod
    def extract_tables(file_path: str) -> List[Dict[str, Any]]:
        """Extract tables from DOCX file."""
        tables = []
        try:
            doc = DocxDocument(file_path)
            for table_num, table in enumerate(doc.tables):
                data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    data.append(row_data)
                
                if data and len(data) > 1:
                    df = pd.DataFrame(data[1:], columns=data[0])
                    tables.append({
                        'table_number': table_num + 1,
                        'data': df.to_dict('records'),
                        'csv': df.to_csv(index=False)
                    })
        except Exception as e:
            print(f"Error extracting tables from DOCX: {e}")
        return tables
    
    @staticmethod
    def extract_comprehensive(file_path: str) -> Dict[str, Any]:
        """Extract text and tables from DOCX."""
        return {
            'text': DocxProcessor.extract_text(file_path),
            'tables': DocxProcessor.extract_tables(file_path)
        }

class ImageProcessor:
    """Process image files to extract text using OCR."""
    
    @staticmethod
    def extract_text(file_path: str) -> str:
        """Extract text from image using OCR."""
        try:
            # Load image
            image = cv2.imread(file_path)
            if image is None:
                return ""
            
            # Convert to RGB
            image_rgb = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
            
            # Extract text using pytesseract
            text = pytesseract.image_to_string(image_rgb)
            return text.strip()
        except Exception as e:
            print(f"Error extracting text from image: {e}")
            return ""
    
    @staticmethod
    def preprocess_image(file_path: str) -> str:
        """Preprocess image for better OCR results."""
        try:
            image = cv2.imread(file_path)
            if image is None:
                return ""
            
            # Convert to grayscale
            gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
            
            # Apply thresholding
            _, thresh = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
            
            # Extract text
            text = pytesseract.image_to_string(thresh)
            return text.strip()
        except Exception as e:
            print(f"Error preprocessing image: {e}")
            return ""

class ExcelProcessor:
    """Process Excel files to extract data."""
    
    @staticmethod
    def extract_data(file_path: str) -> Dict[str, Any]:
        """Extract data from Excel file."""
        try:
            # Read all sheets
            xlsx_file = pd.ExcelFile(file_path)
            sheets_data = {}
            
            for sheet_name in xlsx_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                sheets_data[sheet_name] = {
                    'data': df.to_dict('records'),
                    'csv': df.to_csv(index=False),
                    'summary': f"Sheet '{sheet_name}' contains {len(df)} rows and {len(df.columns)} columns"
                }
            
            return sheets_data
        except Exception as e:
            print(f"Error extracting data from Excel: {e}")
            return {}

class TextProcessor:
    """Process plain text files."""
    
    @staticmethod
    def extract_text(file_path: str) -> str:
        """Extract text from plain text file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            # Try with different encoding
            try:
                with open(file_path, 'r', encoding='latin-1') as file:
                    return file.read()
            except Exception as e:
                print(f"Error reading text file: {e}")
                return ""
        except Exception as e:
            print(f"Error reading text file: {e}")
            return ""

class JSONProcessor:
    """Process JSON files."""
    
    @staticmethod
    def extract_data(file_path: str) -> Dict[str, Any]:
        """Extract data from JSON file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                data = json.load(file)
                return {
                    'data': data,
                    'text': json.dumps(data, indent=2),
                    'summary': f"JSON file contains {len(data) if isinstance(data, (list, dict)) else 1} items"
                }
        except Exception as e:
            print(f"Error reading JSON file: {e}")
            return {}

class UniversalFileProcessor:
    """Universal file processor that handles multiple file types."""
    
    def __init__(self):
        self.processors = {
            'pdf': PDFProcessor,
            'docx': DocxProcessor,
            'doc': DocxProcessor,
            'txt': TextProcessor,
            'csv': self._process_csv,
            'xlsx': ExcelProcessor,
            'xls': ExcelProcessor,
            'json': JSONProcessor,
            'jpg': ImageProcessor,
            'jpeg': ImageProcessor,
            'png': ImageProcessor,
            'bmp': ImageProcessor,
            'tiff': ImageProcessor,
        }
    
    def process_file(self, file_path: str) -> Dict[str, Any]:
        """Process file based on its extension."""
        ext = os.path.splitext(file_path)[1].lower().lstrip('.')
        
        if ext not in self.processors:
            return {'error': f'Unsupported file type: {ext}'}
        
        try:
            if ext in ['pdf', 'docx', 'doc']:
                processor = self.processors[ext]
                return processor.extract_comprehensive(file_path)
            elif ext in ['xlsx', 'xls']:
                processor = self.processors[ext]
                return processor.extract_data(file_path)
            elif ext in ['jpg', 'jpeg', 'png', 'bmp', 'tiff']:
                processor = self.processors[ext]
                return {
                    'text': processor.extract_text(file_path),
                    'preprocessed_text': processor.preprocess_image(file_path)
                }
            elif ext == 'txt':
                processor = self.processors[ext]
                return {'text': processor.extract_text(file_path)}
            elif ext == 'json':
                processor = self.processors[ext]
                return processor.extract_data(file_path)
            elif ext == 'csv':
                return self._process_csv(file_path)
            else:
                return {'error': f'No processor implemented for: {ext}'}
        
        except Exception as e:
            return {'error': f'Error processing file: {str(e)}'}
    
    def _process_csv(self, file_path: str) -> Dict[str, Any]:
        """Process CSV file."""
        try:
            df = pd.read_csv(file_path)
            return {
                'data': df.to_dict('records'),
                'csv': df.to_csv(index=False),
                'summary': f"CSV file contains {len(df)} rows and {len(df.columns)} columns"
            }
        except Exception as e:
            return {'error': f'Error processing CSV: {str(e)}'}
    
    def chunk_text(self, text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
        """Split text into overlapping chunks."""
        if not text:
            return []
        
        chunks = []
        start = 0
        
        while start < len(text):
            end = min(start + chunk_size, len(text))
            chunk = text[start:end]
            chunks.append(chunk)
            
            if end == len(text):
                break
            
            start = end - overlap
        
        return chunks